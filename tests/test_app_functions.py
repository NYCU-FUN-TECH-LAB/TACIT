"""
app.py 裡不需要介面就能測的函式。

為什麼這一支存在
----------------
app.py 是介面檔，但不是每一行都是「UI 層、無法自動化」。
裡面有四段**根本不碰 Streamlit**：

    read_docx               24 行   讀取上傳的逐字稿（每個人都會走這條）
    fix_newlines_in_strings 15 行   修復模型輸出 JSON 裡的裸換行
    save_record             13 行   寫回磁碟——provenance 在這裡保住或丟掉
    build_system_prompt     27 行   由作用中框架產生編碼提示詞

測試不能只測到模組邊界就停。`save_record` 尤其如此：紀錄的 provenance
會不會在存檔時掉，取決於這個函式，只在 schema 層測是不夠的。

匯入 app 會觸發 Streamlit 的 bare mode 警告，那是預期的——這一支只呼叫
純函式，不進入任何 st.* 的繪製路徑。
"""
import io
import json
import os
import sys
import tempfile

import docx

import tacit_framework as F
import tacit_schema as S
import app as APP

FAIL = []


def ok(label, cond, note=""):
    print(("  PASS  " if cond else "  FAIL  ") + label + (f"  ({note})" if note else ""))
    if not cond:
        FAIL.append(label)


def eq(label, got, want):
    ok(label, got == want, f"got={got!r} want={want!r}" if got != want else "")


print("=" * 70)
print("測試 1：read_docx —— 段落式與表格式版面都要讀得出來")
print("=" * 70)
# make_demo_data 會刻意產出兩種版面（部分逐字稿是表格），因為真實世界的
# 逐字稿兩種都有。只支援其中一種的話，上傳自己檔案的人會拿到空白。


def para_doc(lines):
    d = docx.Document()
    for ln in lines:
        d.add_paragraph(ln)
    b = io.BytesIO()
    d.save(b)
    b.seek(0)
    return b


def table_doc(headers, rows, title=None):
    d = docx.Document()
    if title:
        d.add_paragraph(title)
    t = d.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = v
    b = io.BytesIO()
    d.save(b)
    b.seek(0)
    return b


txt = APP.read_docx(para_doc(["Interviewer: 開場問題", "", "   ", "R: 這是回答"]))
ok("段落式讀得出來", "開場問題" in txt and "這是回答" in txt, repr(txt))
ok("空白段落被略過", "\n\n" not in txt, repr(txt))

txt = APP.read_docx(table_doc(["Speaker", "Content"],
                              [["Interviewer", "第一個問題"],
                               ["R01", "第一個回答"],
                               ["", "沒有講者的一列"]]))
ok("表格式讀得出來", "第一個回答" in txt, repr(txt))
ok("講者被標成【】", "【R01】第一個回答" in txt, repr(txt))
ok("沒有講者時不產生空的【】", "【】" not in txt, repr(txt))
ok("表頭本身不會被當成內容", "Speaker" not in txt, repr(txt))

# 中文欄名也要認得——使用者的逐字稿多半是中文表頭
txt = APP.read_docx(table_doc(["講者", "內容"], [["受訪者A", "中文表頭的內容"]]))
ok("中文表頭認得出來", "【受訪者A】中文表頭的內容" in txt, repr(txt))

# 不是逐字稿的表格（例如附在文件裡的資料表）走通用路徑，不可整張丟掉
txt = APP.read_docx(table_doc(["項目", "數值"], [["受訪次數", "3"]]))
ok("非逐字稿表格改用通用串接", "受訪次數 | 3" in txt, repr(txt))

eq("空白文件回空字串", APP.read_docx(para_doc([])), "")

print()
print("=" * 70)
print("測試 2：fix_newlines_in_strings —— 模型輸出裡的裸換行")
print("=" * 70)
# 地端模型很常在 JSON 字串值裡直接放換行，那是不合法的 JSON。這個函式
# 只轉義「字串內」的換行，字串外的排版換行必須原樣保留，否則會把
# 本來合法的 JSON 弄壞。
raw = '{"quote": "第一行\n第二行", "n": 1}'
fixed = APP.fix_newlines_in_strings(raw)
try:
    d = json.loads(fixed)
    ok("修復後可以解析", True)
    eq("換行保留在值裡", d["quote"], "第一行\n第二行")
    eq("其他欄位不受影響", d["n"], 1)
except json.JSONDecodeError as e:
    ok("修復後可以解析", False, str(e))

pretty = '{\n  "a": 1,\n  "b": 2\n}'
eq("字串外的排版換行原樣保留", APP.fix_newlines_in_strings(pretty), pretty)

esc = '{"a": "已經跳脫過的\\n不要重複處理"}'
eq("已跳脫的序列不重複處理", APP.fix_newlines_in_strings(esc), esc)

tab = '{"a": "有\t定位字元"}'
ok("字串內的定位字元也轉義", json.loads(APP.fix_newlines_in_strings(tab))["a"] == "有\t定位字元")

quoted = '{"a": "含有 \\" 逸出引號\n的值"}'
try:
    json.loads(APP.fix_newlines_in_strings(quoted))
    ok("逸出引號不會讓字串邊界判斷錯亂", True)
except json.JSONDecodeError as e:
    ok("逸出引號不會讓字串邊界判斷錯亂", False, str(e))

print()
print("=" * 70)
print("測試 3：save_record —— provenance 必須真的落到磁碟")
print("=" * 70)
# 要防的失效：直接把 _meta 賦成新的兩個鍵，
# 於是每存一次檔，「這份紀錄是哪個模型、哪個端點跑出來的」就消失一次。
# 而複核正是最需要留下軌跡的動作。schema 層有測，但**這個函式本身**
# 也要被執行到——沒測到的地方最容易復發。
_saved_dir = APP.SAVE_DIR
_tmp = tempfile.mkdtemp()
try:
    APP.SAVE_DIR = _tmp
    F.activate_by_id("ri_stilgoe_2013")
    DESC = "ollama/llama3.1:8b-instruct-q4_K_M@http://localhost:11434"

    rec = S.migrate_record({
        S.RESPONDENT: "I01 (Engineer)",
        S.SEGMENTS: [{S.SEGMENT_ID: "S001", S.TITLE: "t", S.QUOTE: "q",
                      S.FULL_TEXT: "f", S.CODES_F: [
                          {S.DIMENSION: S.ANTICIPATION, S.POLARITY: "P",
                           S.RATIONALE: "r"}]}],
        S.META: {"source": DESC, "coded_at": "2026-08-01T10:00:00",
                 "framework_id": "ri_stilgoe_2013"},
    })
    fn = APP.save_record(rec)
    body = json.load(open(os.path.join(_tmp, fn), encoding="utf-8"))

    ok("檔案有寫出來", os.path.isfile(os.path.join(_tmp, fn)), fn)
    eq("端點描述留在磁碟上", body[S.META].get("source"), DESC)
    eq("編碼時間沒有被蓋掉", body[S.META].get("coded_at"), "2026-08-01T10:00:00")
    eq("框架識別碼有記下來", body[S.META].get("framework_id"), "ri_stilgoe_2013")
    eq("schema 版本有更新", body[S.META].get("schema_version"), S.SCHEMA_VERSION)
    ok("有記下最後存檔時間", bool(body[S.META].get("last_saved")))
    # _meta 是**唯一**應該被寫出去的底線鍵——provenance 就掛在它上面。
    # 其餘底線鍵（_file、_source、_fw_mismatch、_dropped）純屬記憶體內的
    # 標記，寫進檔案只會在下次載入時被當成資料。
    ok("除了 _meta 之外沒有底線鍵被寫進檔案",
       [k for k in body if k.startswith("_")] == [S.META], str(list(body)))
    rec["_file_marker_should_not_persist"] = "x"
    rec["_fw_mismatch"] = ("a", "b")
    body_x = json.load(open(os.path.join(_tmp, APP.save_record(rec)),
                            encoding="utf-8"))
    ok("記憶體內的底線標記不會外洩到檔案",
       "_fw_mismatch" not in body_x and
       "_file_marker_should_not_persist" not in body_x, str(list(body_x)))

    # 再存一次（模擬複核後儲存）——provenance 必須還在。
    # 要防的正是這條路徑：改一個碼、按下儲存，來源欄位就永久消失。
    rec[S.SEGMENTS][0][S.CODES_F][0][S.RATIONALE] = "研究者改過的理由"
    fn2 = APP.save_record(rec)
    body2 = json.load(open(os.path.join(_tmp, fn2), encoding="utf-8"))
    eq("複核後再存，端點描述仍在", body2[S.META].get("source"), DESC)
    eq("複核後再存，編碼時間仍在", body2[S.META].get("coded_at"),
       "2026-08-01T10:00:00")
    eq("複核的內容有寫進去",
       body2[S.SEGMENTS][0][S.CODES_F][0][S.RATIONALE], "研究者改過的理由")
    eq("同一筆紀錄寫回同一個檔名（不會每存一次多一個檔）", fn2, fn)

    # 檔名要能安全處理受訪者名稱裡的路徑字元
    bad = S.migrate_record({S.RESPONDENT: 'a/b:c*d?e"f<g>h|i',
                            S.SEGMENTS: [{S.SEGMENT_ID: "S001", S.CODES_F: [
                                {S.DIMENSION: S.ANTICIPATION, S.POLARITY: "P"}]}]})
    fn3 = APP.save_record(bad)
    ok("危險字元不會出現在檔名",
       not any(c in fn3 for c in '/\\:*?"<>|'), fn3)
    ok("該檔案真的存在", os.path.isfile(os.path.join(_tmp, fn3)), fn3)
finally:
    APP.SAVE_DIR = _saved_dir

print()
print("=" * 70)
print("測試 3b：隨附的參考編碼唯讀——複核後另存新檔（#36）")
print("=" * 70)
# 要防的失效：analyses/ 裡的 demo 參考編碼同時是論文 Table 6 的資料來源。
# 在複核頁籤改一個碼並儲存，若直接覆寫原檔，參考標準就悄悄變了：只改一個碼，
# 卡方就從 10.04/.018 變成 10.454/.0151，引用這批語料的數字都對不上。
_saved_dir = APP.SAVE_DIR
_tmp = tempfile.mkdtemp()
try:
    APP.SAVE_DIR = _tmp
    F.activate_by_id("ri_stilgoe_2013")
    ref_fn = "demo_en_Z01.json"
    ref_body = {
        S.RESPONDENT: "Z01 (Test)",
        S.SEGMENTS: [{S.SEGMENT_ID: "S001", S.TITLE: "t", S.QUOTE: "q",
                      S.FULL_TEXT: "f", S.CODES_F: [
                          {S.DIMENSION: S.ANTICIPATION, S.POLARITY: "P",
                           S.RATIONALE: "r"}]}],
        S.META: {"schema_version": S.SCHEMA_VERSION,
                 "framework_id": "ri_stilgoe_2013",
                 "source": "demo/reference-coding"},
    }
    with open(os.path.join(_tmp, ref_fn), "w", encoding="utf-8") as f:
        json.dump(ref_body, f, ensure_ascii=False, indent=2)
    before = open(os.path.join(_tmp, ref_fn), encoding="utf-8").read()

    recs, _ = APP.load_records_from_dir(_tmp, "pro")
    rec = recs[0]
    ok("載入時標記為參考編碼", rec.get("_reference") is True, str(rec.get("_reference")))

    rec[S.SEGMENTS][0][S.CODES_F][0][S.RATIONALE] = "研究者改過"
    new_fn = APP.save_record(rec)
    after = open(os.path.join(_tmp, ref_fn), encoding="utf-8").read()
    ok("原檔一個位元組都沒變", before == after)
    ok("修改存到另一個檔", new_fn != ref_fn and os.path.isfile(os.path.join(_tmp, new_fn)),
       new_fn)
    ok("新檔名看得出來源", new_fn.startswith("demo_en_Z01_reviewed_"), new_fn)
    body = json.load(open(os.path.join(_tmp, new_fn), encoding="utf-8"))
    eq("新檔記下它是從哪一份分出來的", body[S.META].get("derived_from"), ref_fn)
    eq("參考編碼的來源標記保留在 reference_source",
       body[S.META].get("reference_source"), "demo/reference-coding")
    eq("修改的內容在新檔裡", body[S.SEGMENTS][0][S.CODES_F][0][S.RATIONALE], "研究者改過")
    eq("記憶體內的紀錄改指向新檔", rec.get("_file"), new_fn)

    # 再存一次：新檔不是參考編碼，應該原地覆寫，不再分出第三個檔
    fn_again = APP.save_record(rec)
    eq("第二次儲存寫回同一個新檔", fn_again, new_fn)
    eq("目錄裡只有原檔與一個新檔", len([x for x in os.listdir(_tmp) if x.endswith(".json")]), 2)

    # 重新掃描：新檔不能再被當成參考編碼，否則下一次複核又會分出一個
    recs2, _ = APP.load_records_from_dir(_tmp, "pro")
    flags = {r["_file"]: r.get("_reference") for r in recs2}
    eq("原檔仍是參考編碼", flags.get(ref_fn), True)
    eq("複核後的新檔不是參考編碼", flags.get(new_fn), False)
finally:
    APP.SAVE_DIR = _saved_dir

print()
print("=" * 70)
print("測試 4：build_system_prompt —— 提示詞必須由作用中框架產生")
print("=" * 70)
# 提示詞裡不可以有寫死的 RI 維度。主題歸納那一層有自己的測試守著，
# 編碼這一層也要有——沒有測試的正確只是剛好正確。
_probe = {
    "ri_stilgoe_2013": (["anticipation", "reflexivity", "engagement",
                         "responsiveness"],
                        ["performance_expectancy", "effort_expectancy"]),
    "utaut_venkatesh_2003": (["performance_expectancy", "effort_expectancy",
                              "social_influence", "facilitating_conditions"],
                             ["anticipation", "reflexivity", "responsiveness"]),
}
for fid, (mine, theirs) in _probe.items():
    F.activate_by_id(fid)
    fw = F.active()
    p = APP.build_system_prompt(fw, "en", "a short english sample")
    print(f"\n  [{fid}] has_polarity={fw.has_polarity}  prompt={len(p)} chars")
    ok(f"{fid}：自己的每個維度都出現在提示詞裡",
       all(d in p for d in mine),
       str([d for d in mine if d not in p]))
    ok(f"{fid}：不會出現別的框架的維度",
       not any(d in p for d in theirs),
       str([d for d in theirs if d in p]))
    _name = fw.name() if callable(getattr(fw, "name", None)) else (fw.name or "")
    ok(f"{fid}：提示詞裡有框架名稱",
       bool(_name) and _name.split()[0] in p, f"name={_name!r}")

# 語料的**種類**也要跟著框架走，不只是維度與屬性。
#
# 要防的失效：換成永續報告書的框架之後，維度與屬性都正確了，提示詞卻仍寫著
# 「thematic analysis of an interview transcript」，並要模型從「自我介紹、
# 職稱、年資」抽屬性——報告書裡沒有這些東西，模型只能亂填。維度可插拔而
# 語料種類不可插拔，等於論文那句「任何演繹式架構都適用」只成立一半。
F.activate_by_id("ri_stilgoe_2013")
_p_iv = APP.build_system_prompt(F.active(), "en", "sample")
ok("未宣告 corpus 時仍是訪談語彙",
   "interview transcript" in _p_iv and "self-introduction" in _p_iv)
ok("冠詞正確（an interview，不是 a interview）",
   "an interview transcript" in _p_iv and "a interview" not in _p_iv)

_esg = F.load_dict({
    "framework_id": "corpus_probe", "name": {"en": "Corpus probe"},
    "dimensions": [{"id": "d1", "short": "D1", "label": {"en": "D1"}}],
    "corpus": {
        "document": {"en": "corporate sustainability report"},
        "case": {"en": "reporting entity"},
        "descriptor_source": {"en": "the cover and the assurance statement"}},
})
F.set_active(_esg)
_p_esg = APP.build_system_prompt(F.active(), "en", "sample")
ok("宣告後改用該語料的說法",
   "corporate sustainability report" in _p_esg, _p_esg.split("\n")[1][:80])
ok("訪談語彙完全不殘留",
   "interview transcript" not in _p_esg and "self-introduction" not in _p_esg
   and "years of experience" not in _p_esg)
ok("case 名稱換掉了（不是 respondent）",
   "reporting entity" in _p_esg)
ok("屬性來源換成該語料找得到的地方",
   "the cover and the assurance statement" in _p_esg)
ok("摘要指示也用新的 case 名稱",
   "this reporting entity most strongly exhibits" in _p_esg,
   _p_esg[-260:].replace("\n", " ")[-120:])
F.reset()
F.activate_by_id("ri_stilgoe_2013")

# 無極性框架不該要求模型給極性——那個欄位在該框架下沒有意義
F.activate_by_id("utaut_venkatesh_2003")
_fw = F.active()
if not _fw.has_polarity:
    p = APP.build_system_prompt(_fw, "en", "sample")
    ok("無極性框架的提示詞不索取 polarity",
       '"polarity"' not in p.lower(), "有極性欄位就會拿到無意義的值")
else:
    print("  （UTAUT 框架帶極性模型，跳過這一項）")

# 分析語言會影響提示詞裡的語言指示
F.activate_by_id("ri_stilgoe_2013")
_en = APP.build_system_prompt(F.active(), "en", "english sample text")
_zh = APP.build_system_prompt(F.active(), "zh-Hant", "中文樣本文字")
ok("不同分析語言產生不同的提示詞", _en != _zh)
F.reset()

print()
print("=" * 70)
print("測試 5：單筆試跑的解析與「絕不落地」保證")
print("=" * 70)
# 試跑是校準，不是資料。它最重要的性質是**不會寫進任何紀錄**——一旦它
# 悄悄留下東西，研究者就會在不知情的狀況下把校準用的片段混進正式語料。
# 這裡驗的是試跑走的那條解析鏈（擷取 JSON → 修換行 → migrate），
# 以及走完之後存檔資料夾仍然是空的。
import tacit_llm as LLM                                    # noqa: E402

F.activate_by_id("ri_stilgoe_2013")
_dry_raw = (
    'Here is the coding you asked for:\n'
    '{"segments": [{"segment_id": "S001", "title": "Baseline before rollout",\n'
    '  "quote": "六個月的基線期，\n在任何東西被打開之前",\n'
    '  "full_text": "我們談的是六個月的基線期",\n'
    '  "codes": [{"dimension": "Anticipation", "polarity": "P",\n'
    '             "rationale": "把可評估性設計進去"}]}]}\n'
    'Let me know if you need more.'
)
_js = LLM._extract_first_json(_dry_raw)
ok("從夾雜文字的回覆裡挖得出 JSON", bool(_js))
_rec = S.migrate_record(json.loads(APP.fix_newlines_in_strings(_js)))
_segs = _rec.get(S.SEGMENTS) or []
eq("解析出一個段落", len(_segs), 1)
if _segs:
    eq("標題有保留", _segs[0].get(S.TITLE), "Baseline before rollout")
    ok("引文裡的換行沒有弄壞 JSON", "\n" in _segs[0].get(S.QUOTE, ""),
       repr(_segs[0].get(S.QUOTE)))
    # 大寫的維度名——試跑用的是同一條正規化路徑，所以大小寫容錯必須生效
    eq("大寫維度名被正規化", S.codes_of(_segs[0]), ["ANT-P"])
    eq("理由有保留",
       _segs[0][S.CODES_F][0].get(S.RATIONALE), "把可評估性設計進去")

_tmp2 = tempfile.mkdtemp()
_saved_dir = APP.SAVE_DIR
try:
    APP.SAVE_DIR = _tmp2
    # 完整重走一次試跑的解析鏈，中間不呼叫 save_record
    js = LLM._extract_first_json(_dry_raw)
    S.migrate_record(json.loads(APP.fix_newlines_in_strings(js)))
    ok("試跑走完之後，存檔資料夾仍然是空的",
       os.listdir(_tmp2) == [], str(os.listdir(_tmp2)))
finally:
    APP.SAVE_DIR = _saved_dir

# 模型回了完全不相干的東西時，試跑必須拿到「零段落」而不是崩掉
_junk = S.migrate_record(json.loads('{"segments": []}'))
eq("空結果解析得出來且為零段落", len(_junk.get(S.SEGMENTS) or []), 0)
ok("模型回非 JSON 時擷取器回空字串而不是拋錯",
   LLM._extract_first_json("I cannot help with that.") == "")
F.reset()

print()
print("=" * 70)
print("結果：全部通過 ✅" if not FAIL else f"結果：{len(FAIL)} 項失敗 ❌")
for f in FAIL:
    print(f"  - {f}")
print("=" * 70)
sys.exit(1 if FAIL else 0)
