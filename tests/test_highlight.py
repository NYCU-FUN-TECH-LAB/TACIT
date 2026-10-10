"""驗證 tacit_highlight.py：把編碼標回原始逐字稿，輸出螢光筆標記的 Word 檔。"""
import copy
import io
import sys
import zipfile

from docx import Document

import tacit_framework as F
import tacit_schema as S
import tacit_review as RV
import tacit_highlight as H
import tacit_irr as IRR

FAIL = []


def check(label, got, want):
    ok = got == want
    print(("  PASS  " if ok else "  FAIL  ") + f"{label}: got={got!r} want={want!r}")
    if not ok:
        FAIL.append(label)


def check_true(label, cond, note=""):
    print(("  PASS  " if cond else "  FAIL  ") + label + (f"  ({note})" if note else ""))
    if not cond:
        FAIL.append(label)


F.activate_by_id("ri_stilgoe_2013")

LINES = [
    "Interview transcript P01 — Project director",
    "[Interviewer] Was there any discussion early on about unintended consequences?",
    "[Respondent] We ran a baseline period of six months before anything was switched on. "
    "Honestly we never asked residents until the complaints came in.",
    "[Interviewer] And afterwards?",
    "[Respondent] After the complaints we changed the crossing times within a month.",
]
TRANSCRIPT = "\n".join(LINES)

Q1 = "We ran a baseline period of six months before anything was switched on."
Q2 = "Honestly we never asked residents until the complaints came in."
Q3 = "After the complaints we changed the crossing times within a month."
# 模型把換行寫成空格、把連續空白縮成一個的引文
Q_SPACED = "switched on.  Honestly we never asked residents"


def seg(sid, text, codes):
    return {S.SEGMENT_ID: sid, S.TITLE: sid, S.QUOTE: text, S.FULL_TEXT: text,
            S.CODES_F: [S.make_code(d, p) for d, p in codes]}


def make_record():
    rec = {S.RESPONDENT: "P01 (Project director)", S.DESCRIPTORS: S.blank_descriptors(),
           S.SUMMARY: "", S.TRANSCRIPT: TRANSCRIPT,
           S.SEGMENTS: [seg("S001", Q1, [(S.ANTICIPATION, "P")]),
                        seg("S002", Q2, [(S.ENGAGEMENT, "N"), (S.REFLEXIVITY, "P")]),
                        seg("S003", Q3, [(S.RESPONSIVENESS, "P")]),
                        seg("S004", "This sentence is nowhere in the transcript at all.",
                            [(S.ANTICIPATION, "N")])],
           S.DELETED_SEGMENTS: [], S.META: {"endpoint": "ollama/test@http://localhost:11434"}}
    RV.ensure_all([rec])
    return rec


def highlighted(docx_bytes):
    """Word 檔裡被螢光筆標到的文字與顏色：[(text, colour)]，相鄰同色的 run 併起來。"""
    out = []
    for p in Document(io.BytesIO(docx_bytes)).paragraphs:
        cur = None
        for r in p.runs:
            c = r.font.highlight_color
            if c is None:
                cur = None
                continue
            if cur is not None and out and out[-1][1] == c:
                out[-1] = (out[-1][0] + r.text, c)
            else:
                out.append((r.text, c))
            cur = c
    return out


def body_text(docx_bytes):
    return "\n".join(p.text for p in Document(io.BytesIO(docx_bytes)).paragraphs)


print("=" * 70)
print("測試 1：引文找回原文的位置")
print("=" * 70)
check("逐字相同的引文", H.locate(TRANSCRIPT, Q3), (TRANSCRIPT.index(Q3), TRANSCRIPT.index(Q3) + len(Q3)))
sp = H.locate(TRANSCRIPT, Q_SPACED)
check_true("空白不同的引文仍找得到", sp is not None and TRANSCRIPT[sp[0]:sp[1]].startswith("switched on."))
check("不在原文裡的引文找不到", H.locate(TRANSCRIPT, "not in the transcript"), None)
check("空引文", H.locate(TRANSCRIPT, "   "), None)

print()
print("=" * 70)
print("測試 2：目前的編碼")
print("=" * 70)
rec = make_record()
placed, missing = H.place_marks(rec, H.SOURCE_CURRENT)
check("找到位置的段落", [m["id"] for m in placed], ["S001", "S002", "S003"])
check("找不到位置的段落", [m["id"] for m in missing], ["S004"])
data = H.build_docx(rec, H.SOURCE_CURRENT, lang="en", list_unlocated=True)
hl = highlighted(data)
texts = [t for t, _ in hl]
check_true("三段引文都被螢光筆標出來", all(q in texts for q in (Q1, Q2, Q3)), str(texts)[:200])
colours = H.colour_map([rec])
check("S001 用 anticipation 的顏色", dict(hl)[Q1], colours[S.ANTICIPATION])
check("多重編碼的段落取第一個碼的維度顏色", dict(hl)[Q2],
      colours[S.split_code(S.codes_of(rec[S.SEGMENTS][1])[0])[0]])
check_true("四個維度四種顏色", len({colours[d] for d in F.active().dimensions}) == 4)
txt = body_text(data)
check_true("句尾括號帶段落編號與碼", "[S001 ANT-P]" in txt and "[S003 RES-P]" in txt, "")
check_true("多重編碼的括號列出兩個碼", "[S002 " in txt and "ENG-N" in txt and "REF-P" in txt)
check_true("沒被標到的話原樣保留", "[Interviewer] And afterwards?" in txt)
check_true("要求時，找不到位置的段落列在文末", "nowhere in the transcript" in txt)
check_true("預設不列", "nowhere in the transcript" not in body_text(H.build_docx(rec, H.SOURCE_CURRENT, lang="en")))
check("標點與大小寫不同的引文仍找得到",
      H.locate(TRANSCRIPT, "we ran a baseline period, of six months before anything was switched on"),
      (TRANSCRIPT.index(Q1), TRANSCRIPT.index(Q1) + len(Q1) - 1))
check_true("訪員的提問沒有螢光筆", not any("unintended consequences" in t for t in texts))
# 拿掉括號之後，逐字稿每一行都完整還在
doc = Document(io.BytesIO(data))
plain = ["".join(r.text for r in p.runs if not r.text.startswith(" [S")) for p in doc.paragraphs]
check_true("逐字稿每一行都原樣出現", all(line in plain for line in LINES))

print()
print("=" * 70)
print("測試 3：模型的原始草稿——刪掉的也標，人工新增的不標，碼用複核前的")
print("=" * 70)
rec = make_record()
RV.update_codes(rec[S.SEGMENTS][0], ["ANT-N"], reviewer="t")          # S001 的碼被改掉
RV.delete_segment(rec, "S003", reason="not responsiveness", reviewer="t")
RV.add_segment(rec, "[Interviewer] And afterwards?", ["REF-N"], reviewer="t")
cur = H.build_docx(rec, H.SOURCE_CURRENT, lang="en")
mod = H.build_docx(rec, H.SOURCE_MODEL, lang="en")
ct, mt = body_text(cur), body_text(mod)
check_true("目前版本用改過的碼", "[S001 ANT-N]" in ct and "[S001 ANT-P]" not in ct)
check_true("模型版本用原始的碼", "[S001 ANT-P]" in mt and "[S001 ANT-N]" not in mt)
check_true("刪掉的段落：目前版本不標", Q3 not in [t for t, _ in highlighted(cur)])
check_true("刪掉的段落：模型版本照標", Q3 in [t for t, _ in highlighted(mod)])
check_true("人工新增的段落：目前版本有標",
           any("And afterwards?" in t for t, _ in highlighted(cur)))
check_true("人工新增的段落：模型版本不標",
           not any("And afterwards?" in t for t, _ in highlighted(mod)))
check_true("模型版本寫出模型端點", "ollama/test@http://localhost:11434" in mt)
check_true("沒有複核欄位的紀錄也能出模型版本",
           Q1 in [t for t, _ in highlighted(H.build_docx(
               {**make_record(), S.SEGMENTS: [seg("S001", Q1, [(S.ANTICIPATION, "P")])]},
               H.SOURCE_MODEL, lang="en"))])

print()
print("=" * 70)
print("測試 4：空白版（給人工螢光筆用）")
print("=" * 70)
rec = make_record()
blank = H.build_docx(rec, H.SOURCE_BLANK, lang="en")
bh = highlighted(blank)
check("人工編碼版沒有任何螢光筆", len(bh), 0)
check_true("逐字稿本身沒有任何螢光筆", not any(Q1 in t or Q2 in t for t, _ in bh))
check_true("逐字稿內容完整", all(line in body_text(blank) for line in LINES))
check_true("空白版附上定義、指標與排除條件",
           all(k in body_text(blank) for k in ("Definition", "Indicators", "Not this dimension")))

print()
print("=" * 70)
print("測試 5：重疊的段落、沒有逐字稿、zip")
print("=" * 70)
rec = make_record()
rec[S.SEGMENTS] = [seg("S001", Q1 + " " + Q2, [(S.ANTICIPATION, "P")]),
                   seg("S002", Q2, [(S.ENGAGEMENT, "N")])]
ov = H.build_docx(rec, H.SOURCE_CURRENT, lang="en")
joined = "".join(t for t, _ in highlighted(ov))
check_true("重疊的兩段：文字只出現一次，兩個括號都在",
           joined.count("Honestly we never asked residents") == 1
           and "[S001 ANT-P]" in body_text(ov) and "[S002 ENG-N]" in body_text(ov))
empty = copy.deepcopy(make_record()); empty[S.TRANSCRIPT] = ""
check_true("沒有逐字稿的紀錄不例外，段落列在文末",
           "holds no transcript" in body_text(H.build_docx(empty, H.SOURCE_CURRENT, lang="en")))
check_true("中文介面的說明", "顏色圖例" in body_text(H.build_docx(make_record(), H.SOURCE_CURRENT, lang="zh")))
a, b = make_record(), make_record()
z = zipfile.ZipFile(io.BytesIO(H.build_zip([a, b], H.SOURCE_MODEL, lang="en")))
check("同名受訪者各有一個檔", len(z.namelist()), 2)
check_true("檔名帶來源、不含不能當檔名的字元",
           all(n.endswith("_model.docx") and "(" not in n.split("_model")[0][-1:] and " " not in n
               for n in z.namelist()), str(z.namelist()))

print()
print("=" * 70)
print("測試 6：開放編碼的紀錄（還沒有框架）")
print("=" * 70)
open_rec = {S.RESPONDENT: "W01", S.TRANSCRIPT: TRANSCRIPT,
            "open_segments": [{S.SEGMENT_ID: "O1", S.QUOTE: Q1, S.FULL_TEXT: Q1,
                               "open_codes": ["baseline before rollout"]},
                              {S.SEGMENT_ID: "O2", S.QUOTE: Q3, S.FULL_TEXT: Q3,
                               "open_codes": [{"label": "change after complaints"}]}]}
od = H.build_docx(open_rec, H.SOURCE_CURRENT, lang="en")
oh = dict(highlighted(od))
check_true("開放編碼的兩段都標出來", Q1 in oh and Q3 in oh)
check_true("不同的開放碼不同顏色", oh[Q1] != oh[Q3])
check_true("括號裡是開放碼的標籤", "baseline before rollout" in body_text(od))

print()
print("=" * 70)
print("測試 7：匯出的檔案讀得回來（註解帶碼）")
print("=" * 70)
import tacit_i18n as I
rec = make_record()
placed, _missing = H.place_marks(rec, H.SOURCE_MODEL)
back = H.read_marked_docx(io.BytesIO(H.build_docx(rec, H.SOURCE_MODEL, lang="en")))
check("逐字稿原樣讀回（標題、圖例、句尾括號、文末清單都不算）", back["transcript"], TRANSCRIPT)
check("檔案標題是受訪者", back["respondent"], "P01 (Project director)")
check("每個段落一個註解，位置相同",
      [(m["start"], m["end"]) for m in back["marks"]], [(m["start"], m["end"]) for m in placed])
check("註解裡的碼讀得回來", [sorted(m["codes"]) for m in back["marks"]],
      [sorted(m["codes"]) for m in placed])
check("註解掛著的文字就是那段引文", [m["text"] for m in back["marks"]], [Q1, Q2, Q3])
check("有註解的螢光筆不算「只畫沒寫」", back["highlight_only"], [])
check_true("註解作者標明來源", all(m["author"] == "TACIT (model draft)" for m in back["marks"]))
blank = H.read_marked_docx(io.BytesIO(H.build_docx(rec, H.SOURCE_BLANK, lang="zh")))
check("空白版讀回來：逐字稿完整、沒有任何標記",
      (blank["transcript"], blank["marks"], blank["highlight_only"]), (TRANSCRIPT, [], []))
check_true("空白版寫了怎麼用註解編碼",
           "註解" in body_text(H.build_docx(rec, H.SOURCE_BLANK, lang="zh"))
           and "ANT-P" in body_text(H.build_docx(rec, H.SOURCE_BLANK, lang="zh")))

print()
print("=" * 70)
print("測試 8：人在 Word 裡加註解的檔案")
print("=" * 70)
from docx.enum.text import WD_COLOR_INDEX


def student_docx(comments):
    """
    模擬學生交回來的檔案：一段話選起來加註解。comments 是
    [(行號, 起, 迄, 註解文字)]，同一行可以有好幾個、可以重疊。
    """
    doc = Document()
    pars = []
    for n, line in enumerate(LINES):
        cuts = sorted({0, len(line)} | {x for c in comments if c[0] == n for x in (c[1], c[2])})
        par = doc.add_paragraph()
        runs = []
        for a, b in zip(cuts, cuts[1:]):
            runs.append((a, b, par.add_run(line[a:b])))
        pars.append(runs)
    for n, a, b, text in comments:
        doc.add_comment([r for x, y, r in pars[n] if x >= a and y <= b], text=text,
                        author="student", initials="s")
    return doc, pars


i1, i2 = LINES[2].index(Q1), LINES[2].index(Q2)
i3 = LINES[4].index(Q3)
doc, pars = student_docx([
    (2, i1, i1 + len(Q1), "ANT-P"),
    (2, i2, i2 + len(Q2), "eng-n；REF-P 沒問居民"),                       # 小寫、全形分號、後面帶說明
    (2, i1, i2 + len(Q2), I.code_label("RES-N", "en")),                    # 與前兩個重疊，用完整標籤寫
    (4, i3, i3 + len(Q3), "not sure about this one"),                      # 認不出碼
])
# 另外畫一段螢光筆但沒加註解
pars[1][0][2].font.highlight_color = WD_COLOR_INDEX.YELLOW
buf = io.BytesIO(); doc.save(buf)
got = H.read_marked_docx(io.BytesIO(buf.getvalue()))
check("沒有分隔線的檔案整份都是逐字稿", got["transcript"], TRANSCRIPT)
by_text = {m["text"]: m for m in got["marks"]}
check("識別字", by_text[Q1]["codes"], ["ANT-P"])
check("小寫、全形分號、後面帶說明", by_text[Q2]["codes"], ["REF-P", "ENG-N"])
check("完整標籤、跨兩句的重疊範圍", by_text[Q1 + " " + Q2]["codes"], ["RES-N"])
check("認不出碼的註解留著但沒有碼", (by_text[Q3]["codes"], by_text[Q3]["comment"]),
      ([], "not sure about this one"))
check("只畫螢光筆沒寫註解的段落另外列出", [h["text"] for h in got["highlight_only"]], [LINES[1]])
check("註解作者", by_text[Q1]["author"], "student")

hrec = H.record_from_marked(got, "P01 (Project director)", coder="student_a")
check("認得出碼的註解才成為段落", len(hrec[S.SEGMENTS]), 3)
check("紀錄標明是人工標記", (hrec[S.META]["source"], hrec[S.META]["endpoint"]),
      ("human-marked/student_a", None))
check("人工段落的原始編碼留空", hrec[S.SEGMENTS][0][S.REVIEW][S.ORIGINAL_CODES], [])
check_true("人工紀錄可以再匯出成螢光筆檔",
           Q1 in [t for t, _ in highlighted(H.build_docx(hrec, H.SOURCE_CURRENT, lang="en"))])

print()
print("=" * 70)
print("測試 9：人工標記與模型編碼的一致性")
print("=" * 70)
rec = make_record()
model_marks, n_missing = H.model_marks_of(rec, H.SOURCE_MODEL)
check("模型有一段的引文找不到位置", n_missing, 1)
same = H.compare_marks(TRANSCRIPT, model_marks, model_marks)
check("兩邊相同：κ = 1", same["pooled"][H.IRR.KAPPA], 1.0)
check("兩邊相同：precision、recall = 1，沒有不同的單元",
      (same["precision"], same["recall"], same["differ"]), (1.0, 1.0, []))
check("抽樣框只有受訪者的發言", same["units"], 2)
human = [{"text": Q1, "codes": ["ANT-P"]}, {"text": Q2, "codes": ["ENG-N", "REF-P"]},
         {"text": Q3, "codes": ["RES-N"]}]                                   # 第三段人標的是 RES-N
cmp = H.compare_marks(TRANSCRIPT, human, model_marks)
check("人標了、模型沒標的格", cmp["pooled"][H.IRR.ONLY_A], 1)
check("模型標了、人沒標的格", cmp["pooled"][H.IRR.ONLY_B], 1)
check("precision 與 recall 以人為參照", (cmp["precision"], cmp["recall"]), (0.75, 0.75))
check("不同的單元列得出來", [(d["human"], d["model"]) for d in cmp["differ"]],
      [(["RES-N"], ["RES-P"])])
check_true("太短而放不上任何單元的標記有計數",
           H.compare_marks(TRANSCRIPT, [{"text": "a month", "codes": ["RES-P"]}], model_marks)
           ["human_not_placed"] == 1)

print()
print("=" * 70)
print("測試 10：檔案對應到哪一筆紀錄；無極性框架的碼")
print("=" * 70)
other = make_record(); other[S.RESPONDENT] = "G07 (Clerk)"; other[S.TRANSCRIPT] = "Something else entirely."
check("逐字稿相同的優先", H.match_record(got, [other, rec]), 1)
check("逐字稿對不上時看檔名裡的受訪者編號",
      H.match_record({"transcript": "edited text", "respondent": ""}, [other, rec], "P01_student_a.docx"), 1)
check("對不上就是 None",
      H.match_record({"transcript": "edited text", "respondent": ""}, [other, rec], "notes.docx"), None)
F.activate_by_id("utaut_venkatesh_2003")
check("無極性框架：識別字要大小寫相符", H.parse_codes("PE and SI"), ["PE", "SI"])
check("一般單字不會被當成碼", H.parse_codes("the type of pe lesson, si"), [])
check("無極性框架：完整標籤", H.parse_codes("performance expectancy"), ["PE"])
F.activate_by_id("ri_stilgoe_2013")

print()
print("=" * 70)
print("測試 11：裁決——兩組不同的單元由人定稿")
print("=" * 70)
rec = make_record()
mm, _ = H.model_marks_of(rec, H.SOURCE_MODEL)
hm = [{"text": Q1, "codes": ["ANT-P"]}, {"text": Q2, "codes": ["ENG-N", "REF-P"]}, {"text": Q3, "codes": ["RES-N"]}]
cmp = H.compare_marks(TRANSCRIPT, hm, mm)
check("compare_marks 回傳抽樣框的單元文字", [r["text"][:20] for r in cmp["frame"]], [LINES[2][13:33], LINES[4][13:33]])
diff_units = [d["unit"] for d in cmp["differ"]]
check("只有一個單元不同", len(diff_units), 1)
u = diff_units[0]
for choice, want in ((H.DECISION_A, ["RES-N"]), (H.DECISION_B, ["RES-P"]), (H.DECISION_BOTH, ["RES-N", "RES-P"]),
                     (H.DECISION_NEITHER, None), (H.DECISION_CUSTOM, ["ANT-N"])):
    adj = H.adjudicated_record(cmp, {u: {"decision": choice, "codes": ["ANT-N"], "reason": "r"}},
                               "P01", "judge", "hand", "model", transcript=TRANSCRIPT)
    got = [sorted(S.codes_of(sg)) for sg in adj[S.SEGMENTS] if sg[S.FULL_TEXT] == Q3]
    check(f"裁決 {choice}", got, [want] if want else [])
adj = H.adjudicated_record(cmp, {}, "P01", "judge", transcript=TRANSCRIPT)
check("沒裁決的單元記在 undecided", adj[S.META]["adjudication"]["undecided"], [u])
check("相同的單元直接採用", [sorted(S.codes_of(sg)) for sg in adj[S.SEGMENTS]], [["ANT-P", "ENG-N", "REF-P"]])
adj = H.adjudicated_record(cmp, {u: {"decision": H.DECISION_B, "reason": "model read it as adaptive"}}, "P01", "judge", "hand", "model", transcript=TRANSCRIPT)
h = [sg for sg in adj[S.SEGMENTS] if sg[S.FULL_TEXT] == Q3][0][S.REVIEW][S.HISTORY][0]
check("歷史記下兩組原本的碼、決定、理由與裁決者",
      (h["hand"], h["model"], h["decision"], h["reason"], h["by"]), (["RES-N"], ["RES-P"], "b", "model read it as adaptive", "judge"))
check("定稿紀錄標明來源", adj[S.META]["source"], "adjudicated/judge")
check("裁決計數", adj[S.META]["adjudication"]["counts"]["b"], 1)
check_true("定稿紀錄可以再匯出成螢光筆檔",
           Q3 in [t for t, _ in highlighted(H.build_docx(adj, H.SOURCE_CURRENT, lang="en"))])

print()
print("=" * 70)
print("測試 12：維度層一致性、比對工作簿、並排的兩個螢光筆版本")
print("=" * 70)
dimk = H.dimension_agreement(cmp)
check("維度層的格數 = 單元 × 維度", dimk["n"], cmp["units"] * 4)
check("第三單元 RES-N 對 RES-P 在維度層算一致", dimk["a_only"] + dimk["b_only"], 0)
wb = H.comparison_workbook([("f1.docx", cmp)], "hand", "model")
import openpyxl
sheets = openpyxl.load_workbook(io.BytesIO(wb)).sheetnames
check("工作簿三張表", sheets, ["summary", "per_code", "units"])
sbs = zipfile.ZipFile(io.BytesIO(H.side_by_side_zip(
    {"transcript": TRANSCRIPT, "marks": [{"start": 0, "end": 1, "text": Q1, "comment": "ANT-P", "author": "s", "codes": ["ANT-P"]}],
     "highlight_only": [], "respondent": ""}, rec, H.SOURCE_MODEL, "P01", "student", "en")))
check("並排 zip 有兩個檔", len(sbs.namelist()), 2)

print()
print("=" * 70)
print("測試 13：本文裡的括號標記（【註解：碼; 理由】）當成標記讀，讀完從逐字稿移除")
print("=" * 70)
_d13 = Document()
_d13.add_paragraph("Interview transcript P02")
_p13 = _d13.add_paragraph("[Respondent] We ran a baseline period of six months. 【註解：ANT-P; 先做基線】")
_p13.add_run().add_break()
_p13.add_run("Honestly we never asked residents until the complaints came in. [[ENG-N]] After that we changed the crossing times. 【RES-P; 改了】")
_d13.add_paragraph("[Interviewer] (Chair) And afterwards? [Respondent]")
_b13 = io.BytesIO(); _d13.save(_b13); _b13.seek(0)
_r13 = H.read_marked_docx(_b13)
check("三個括號都讀成標記", [m["codes"] for m in _r13["marks"]], [["ANT-P"], ["ENG-N"], ["RES-P"]])
check("第一個標記涵蓋講者標籤之後、括號之前的那句", _r13["marks"][0]["text"], "We ran a baseline period of six months.")
check("軟換行之後的第二個標記只涵蓋它那一行到括號", _r13["marks"][1]["text"], "Honestly we never asked residents until the complaints came in.")
check("同一行第二個括號從上一個括號之後算起", _r13["marks"][2]["text"], "After that we changed the crossing times.")
check("理由留在 comment 欄", _r13["marks"][0]["comment"], "ANT-P; 先做基線")
check_true("括號已從逐字稿移除", "【" not in _r13["transcript"] and "[[" not in _r13["transcript"])
check_true("講者標籤不是標記、也沒被移除", "[Interviewer] (Chair)" in _r13["transcript"] and "[Respondent]" in _r13["transcript"])
check("標記位置對得上移除括號後的逐字稿", _r13["transcript"][_r13["marks"][1]["start"]:_r13["marks"][1]["end"]], _r13["marks"][1]["text"])
check("軟換行成為換行", _r13["transcript"].count("\n") >= 3, True)
_cmp13 = H.compare_marks(_r13["transcript"], _r13["marks"], _r13["marks"])
check("括號標記可以直接進比對", (_cmp13["units"] > 0, _cmp13["pooled"][IRR.KAPPA]), (True, 1.0))

print()
print("=" * 70)
print("測試 14：只有逐字稿的版本，與編碼者自己的名字")
print("=" * 70)
rec = make_record()
pl = H.build_docx(rec, H.SOURCE_MODEL, lang="zh", author="Hank", plain=True)
pt = body_text(pl)
check("沒有任何螢光筆", highlighted(pl), [])
check_true("沒有圖例、說明、分隔線、段落後括號",
           not any(x in pt for x in ("顏色圖例", "每段螢光筆", "標出", H.SEPARATOR, "[S00")), pt[:120])
check("檔案的段落就是原稿的行，一行不多一行不少",
      [p.text for p in Document(io.BytesIO(pl)).paragraphs], TRANSCRIPT.split(chr(10)))
pd_ = Document(io.BytesIO(pl))
check("註解作者是編碼者的名字", sorted({c.author for c in pd_.comments}), ["Hank"])
check("註解縮寫", sorted({c.initials for c in pd_.comments}), ["H"])
check("檔案屬性的作者與最後修改者", (pd_.core_properties.author, pd_.core_properties.last_modified_by), ("Hank", "Hank"))
check("沒有指定名稱時維持 TACIT 的標示",
      sorted({c.author for c in Document(io.BytesIO(H.build_docx(rec, H.SOURCE_MODEL, lang="zh"))).comments}),
      ["TACIT (model draft)"])
pback = H.read_marked_docx(io.BytesIO(pl))
check("只有逐字稿的檔案讀得回來，文字與原稿相同", pback["transcript"], TRANSCRIPT)
check("標記與有標題的版本相同", [(m["start"], m["end"], sorted(m["codes"])) for m in pback["marks"]],
      [(m["start"], m["end"], sorted(m["codes"])) for m in back["marks"]])
check_true("作者讀成 Hank", all(m["author"] == "Hank" for m in pback["marks"]))
pz = zipfile.ZipFile(io.BytesIO(H.build_zip([rec], H.SOURCE_MODEL, lang="zh", author="Hank", plain=True)))
check("檔名結尾用編碼者的名字", [n.endswith("_Hank.docx") for n in pz.namelist()], [True])
bl = body_text(H.build_docx(rec, H.SOURCE_BLANK, lang="en", plain=True))
check_true("未標記版不受 plain 影響，仍附說明與編碼簿", "Codebook" in bl)

print()
print("=" * 70)
print(f"結果：{'全部通過 ✅' if not FAIL else '失敗項目 ' + str(FAIL)}")
print("=" * 70)
sys.exit(1 if FAIL else 0)
