"""
bench_models.py — 各模型在本工具上的實測比較
=============================================

    python bench_models.py --provider ollama --models llama3.1:8b,llama3.3:70b
    python bench_models.py --provider gemini --models gemini-3.6-flash --api-key ...
    python bench_models.py --probe-only --models a,b,c        只跑格式探測（幾秒）

為什麼要有這支
--------------
主題歸納那一步會因為模型不同而整段失效，而且失效的方式不是「品質略差」：
一次實測跑出 9 個主題，其中 7 個掉進未歸屬，四個維度有三個掛零——但模型
自己寫的歸屬理由明明指名了 anticipation 與 responsiveness。它不是判斷得
不好，是沒照輸出約定回答。

工具不會替使用者擋掉任何模型（那是研究者的判斷，而且寫死的模型名單一定
會過期）。工具要做的是**把差別量出來**，讓「為什麼留這幾個選項」有依據。

這支腳本的產出就是那份依據：一張模型 × 任務的表，可以直接放進論文。

量什麼
------
  contract    輸出約定遵循度（0–1）。幾秒鐘就跑得完，最該先跑的一項
  coding      分窗編碼：段落數、每萬字元密度、編碼數、逐字引文比對通過率、
              以缺席為證據而被擋下的碼、標題退化成維度名的段落數、耗時
  themes      主題歸納：主題數、不合格維度值、涵蓋到幾個維度、耗時

  編碼走的是與 app.py 完全相同的分窗路徑（tacit_coding.code_transcript）。
  這一點不能省：整份逐字稿一次送出去正是會讓召回率崩掉的做法——
  一份 46,730 字元的稿件只回 4 段。兩邊不一致的話，
  Table 5 量到的是使用者實際上不會遇到的行為，那比沒有 Table 5 更糟。

不量什麼
--------
**不宣稱哪一顆模型「比較準」。** 參考編碼是與逐字稿一起撰寫的固定比較基準，
不是 ground truth；拿它算出來的一致性只能說明「與該基準的距離」，不能當成
正確率。這個界線在論文裡也要維持。

安全性
------
所有輸出寫進 bench_out/，**不碰 analyses/**。參考編碼是語料的一部分，
基準測試不該覆蓋它。
"""

import argparse
import glob
import json
import os
import re
import sys
import time
from datetime import datetime

import docx

import tacit_coding as CH
import tacit_framework as F
import tacit_schema as S
import tacit_themes as T
import tacit_llm as LLM
import tacit_review as RV

OUT_DIR = "bench_out"
RESULTS = os.path.join(OUT_DIR, "bench_results.json")


# ---------------------------------------------------------------------
def read_docx(path):
    doc = docx.Document(path)
    out = []
    for para in doc.paragraphs:
        if para.text.strip():
            out.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            vals = [c.text.strip() for c in row.cells if c.text.strip()]
            if vals:
                out.append(" | ".join(vals))
    return "\n".join(out)


def build_system_prompt(fw, analysis_lang, sample=""):
    """與 app.py 相同的編碼提示詞。這裡重建一份，避免 import Streamlit。"""
    # 在 streamlit run 以外 import app，每個元件都會印一行
    # "missing ScriptRunContext"——跑完 24 份上千行，真正的錯誤訊息被淹沒。
    # 調 streamlit 的 log level 沒有用（它的 handler 不看），只能在 import 的
    # 那一刻整個關掉 logging，import 完立刻恢復，後面的真正錯誤照常顯示。
    import logging
    logging.disable(logging.WARNING)
    try:
        import app  # noqa: E402  —— bare mode 下可 import，只取純函式
    finally:
        logging.disable(logging.NOTSET)
    return app.build_system_prompt(fw, analysis_lang, sample)


# ---------------------------------------------------------------------
def bench_contract(ep):
    """輸出約定遵循度。最便宜的一項，永遠先跑。"""
    call = LLM.make_callable(ep, temperature=0.0)
    t0 = time.time()
    compliance, detail = T.probe_output_contract(call, F.active().dimensions)
    return {
        "compliance": round(compliance, 3),
        "verdict": T.contract_verdict(compliance, detail),
        "answered": detail["answered"],
        "asked": detail["asked"],
        "valid_json": detail["valid_json"],
        "non_conforming": detail["non_conforming"][:6],
        "error": detail["error"],
        "seconds": round(time.time() - t0, 1),
    }


def bench_coding(ep, corpus, lang, limit=None, window=CH.DEFAULT_WINDOW_CHARS,
                 overlap=CH.DEFAULT_OVERLAP_CHARS, save_dir=None):
    """
    對整批逐字稿跑編碼，回傳計數與耗時。不寫進 analyses/。

    **走的路徑必須與 app.py 相同。** 這支腳本若把整份逐字稿一次送出去，
    而 app.py 是分窗編碼——兩邊不一致的話，Table 5 量的是一個使用者
    實際上不會遇到的行為，而那比沒有 Table 5 更糟。

    量的東西不只段落數、編碼數、逐字率——那三個指標分不出
    「模型有沒有在捏造」：
      absence_codes      以「這段話沒有表現出 X」為理由而被擋下的碼。
                         數字大代表這顆模型會憑沉默生出負向編碼。
      degenerate_titles  標題被填成維度名（"Reflexivity - N"）的段落數。
                         代表模型在照抄輸出範例，而不是在讀文本。
      per_10k            每萬字元的段落數。人工編碼的參考值約 13.9；
                         低到 1 以下代表模型回的是摘錄而不是編碼。
    """
    files = sorted(glob.glob(os.path.join("demo_data", corpus, "*.docx")))
    if limit:
        files = files[:limit]
    fw = F.active()
    n_seg = n_code = n_fail = 0
    quote_ok = quote_total = 0
    n_absence = n_degenerate = n_chars = 0
    n_resumed = n_chunk_errors = 0
    seconds = 0.0
    stopped_on_quota = None
    records = []
    labels = [fw.label(d, "en") for d in fw.dimensions]

    def tally(rec, transcript, drops_absence):
        nonlocal n_seg, n_code, quote_ok, quote_total, n_absence, n_degenerate
        nonlocal n_chars, n_chunk_errors
        n_chars += len(transcript)
        n_absence += drops_absence
        n_degenerate += len(S.degenerate_titles(rec, labels=labels))
        n_chunk_errors += len(((rec.get(S.META) or {}).get("chunk_errors")) or [])
        for seg in rec[S.SEGMENTS]:
            n_seg += 1
            n_code += len(seg[S.CODES_F])
            q = (seg.get(S.QUOTE) or "").strip()
            if q:
                quote_total += 1
                if RV.normalize(q) in RV.normalize(transcript):
                    quote_ok += 1

    for path in files:
        stem = os.path.splitext(os.path.basename(path))[0]
        saved = os.path.join(save_dir, stem + ".json") if save_dir else None

        # 換一組 API key 接著跑：已經完整編完的逐字稿直接沿用，不重送。
        # 只沿用有 _bench 標記的檔案——那個標記只在整份編完、沒有撞到額度
        # 上限時才會寫入，所以讀進來的一定是完整結果。
        if saved and os.path.isfile(saved):
            with open(saved, encoding="utf-8") as f:
                rec = json.load(f)
            bench = rec.get("_bench") or {}
            if bench.get("complete") and bench.get("endpoint") == ep.describe():
                records.append(rec)
                seconds += float(bench.get("seconds") or 0)
                tally(rec, rec.get(S.TRANSCRIPT) or "", int(bench.get("absence") or 0))
                n_resumed += 1
                continue

        transcript = read_docx(path)
        sysmsg = build_system_prompt(fw, lang, transcript)
        drops = []
        t_file = time.time()

        def code_one(text, n, total, _p=path, _s=sysmsg, _d=drops):
            raw = LLM.complete(
                ep, f"File: {os.path.basename(_p)}\n\n"
                    f"Excerpt {n + 1} of {total}:\n{text}",
                system=_s, temperature=0.2, json_mode=True)
            js = LLM._extract_first_json(raw)
            if not js:
                raise ValueError("no JSON in response")
            return S.migrate_record(json.loads(js), dropped=_d)

        try:
            if len(transcript) > window:
                merged = CH.code_transcript(transcript, code_one,
                                            window=window, overlap=overlap)
                merged.pop("_chunk_summaries", None)
                # 分窗編碼會吞掉單一窗口的錯誤繼續跑，這對壞 JSON 是對的；
                # 但額度用完不是模型行為，那份紀錄只編了一半。當成完成存下
                # 來，Table 5 的段落數就會悄悄變少。所以撞到額度就整份不算。
                quota = [e for e in (merged[S.META].get("chunk_errors") or [])
                         if e["error"].startswith((LLM.RateLimited.__name__,
                                                   LLM.NotConfigured.__name__,
                                                   LLM.Unavailable.__name__))]
                if quota:
                    raise LLM.RateLimited(quota[0]["error"])
                if not merged.get(S.SEGMENTS):
                    raise ValueError("no segments after merge")
                rec = S.migrate_record(merged)
            else:
                rec = code_one(transcript, 0, 1)
        # 服務端滿載在自動重試後仍失敗，一樣不是模型行為：停下來，稍後同一行指令接著跑
        except (LLM.RateLimited, LLM.NotConfigured, LLM.Unavailable) as e:
            stopped_on_quota = f"{os.path.basename(path)}: {type(e).__name__}: {e}"
            print(f"\n      ■ 額度用完或無法連線，停在 {os.path.basename(path)}。"
                  f"\n        已完成的 {len(records)} 份已存檔。換一組 API key 後"
                  f"用同一行指令重跑，會從這一份接著跑。\n        {str(e)[:160]}",
                  file=sys.stderr)
            break
        except Exception as e:
            n_fail += 1
            print(f"      ! {os.path.basename(path)}: {type(e).__name__}: {e}",
                  file=sys.stderr)
            continue

        took = time.time() - t_file
        absence = sum(1 for d in drops if d.get("reason") == S.DROP_ABSENCE_RATIONALE)
        rec[S.TRANSCRIPT] = transcript
        rec["_bench"] = {"complete": True, "endpoint": ep.describe(),
                         "provenance": ep.provenance(temperature=0.2,
                                                     window_chars=window,
                                                     overlap_chars=overlap),
                         "seconds": round(took, 1), "absence": absence}
        records.append(rec)
        seconds += took
        # 逐篇存檔：與參考編碼算一致性要用到，額度用完換 key 時也從這裡接續。
        if saved:
            os.makedirs(save_dir, exist_ok=True)
            with open(saved, "w", encoding="utf-8") as f:
                json.dump(rec, f, ensure_ascii=False, indent=1)
        tally(rec, transcript, absence)
        print(f"      ✓ {stem}  {len(rec[S.SEGMENTS])} 段  {took:.0f}s", flush=True)

    return {
        "transcripts": len(files),
        "completed": len(records),
        "resumed": n_resumed,
        "stopped_on_quota": stopped_on_quota,
        "chunk_errors": n_chunk_errors,
        "failed": n_fail,
        "segments": n_seg,
        "codes": n_code,
        "chars": n_chars,
        "per_10k": round(n_seg / n_chars * 10000, 2) if n_chars else None,
        # 逐字引文比對：模型有沒有真的引用原文，而不是改寫或編造
        "verbatim_rate": (round(quote_ok / quote_total, 3)
                          if quote_total else None),
        "absence_codes": n_absence,
        "degenerate_titles": n_degenerate,
        "window_chars": window,
        "overlap_chars": overlap,
        # 各份實際編碼時間的總和。換 key 之間停下來的時間不算進去，
        # 否則同一顆模型分三天跑完，耗時會寫成三天。
        "minutes": round(seconds / 60, 1),
    }, records


def bench_themes(ep, records):
    """主題歸納。這是實測中最會出事的一步。"""
    items = T.collect_first_order(records)
    if len(items) < 10:
        return {"skipped": "too few first-order concepts"}
    call = LLM.make_callable(ep, temperature=0.4)
    t0 = time.time()
    try:
        themes, note, dbg = T.induce_themes(items, call, chunk_size=60,
                                            min_size=2, analysis_lang="en")
    except Exception as e:
        return {"error": f"{type(e).__name__}: {e}",
                "minutes": round((time.time() - t0) / 60, 1)}
    dims = F.active().dimensions
    covered = {d for t in themes if t[S.AGG_DIMENSION] in dims
               for d in [t[S.AGG_DIMENSION]]}
    unassigned = sum(1 for t in themes if t[S.AGG_DIMENSION] == S.UNASSIGNED)
    return {
        "first_order": len(items),
        "themes": len(themes),
        "unassigned": unassigned,
        "dimensions_covered": len(covered),
        "dimensions_total": len(dims),
        "empty_dimensions": [d for d in dims if d not in covered],
        # 這一欄是查問題唯一的線索：模型實際寫了什麼不合格的值
        "unrecognised_dimensions": (dbg.get("unrecognised_dimensions") or [])[:8],
        "note_written": bool(note and note.strip().lower() not in ("", "none")),
        "minutes": round((time.time() - t0) / 60, 1),
    }


# ---------------------------------------------------------------------
def markdown_table(results):
    """兩張表：一張給 Table 5（編碼），一張給 Table 5b（歸納）。"""
    out = []
    out.append("### Coding\n")
    out.append("| Model | Corpus | Transcripts | Segments | Seg/10k | Codes | "
               "Verbatim | Absence | Degen. titles | Minutes |")
    out.append("|---|---|---|---|---|---|---|---|---|---|")
    for r in results:
        for corpus, c in (r.get("coding") or {}).items():
            v = "—" if c.get("verbatim_rate") is None else f"{c['verbatim_rate']:.0%}"
            done_n = c.get("completed", c["transcripts"] - c["failed"])
            shown = (f"{c['transcripts']}" if done_n == c["transcripts"]
                     else f"{done_n}/{c['transcripts']} (incomplete)")
            out.append(f"| {r['model']} | {corpus} | {shown}"
                       f"{' (' + str(c['failed']) + ' failed)' if c['failed'] else ''}"
                       f" | {c['segments']} | {c.get('per_10k', '—')}"
                       f" | {c['codes']} | {v} | {c.get('absence_codes', '—')}"
                       f" | {c.get('degenerate_titles', '—')} | {c['minutes']} |")
    out.append("\nSeg/10k = coded segments per 10,000 characters of transcript. "
               "Manual coding of comparable interview material runs around 13.9. "
               "Absence = codes discarded because the rationale offered the "
               "absence of a stance as evidence for it. Degen. titles = segment "
               "titles filled in with the dimension name instead of a sub-theme "
               "label. Both are model-compliance failures, not quality scores.")
    out.append("\n### Theme induction\n")
    out.append("| Model | Contract | Themes | Unassigned | Dimensions covered | Minutes |")
    out.append("|---|---|---|---|---|---|")
    for r in results:
        th = r.get("themes") or {}
        ct = r.get("contract") or {}
        if th.get("skipped") or th.get("error"):
            out.append(f"| {r['model']} | {ct.get('compliance','—')} | "
                       f"{th.get('error') or th.get('skipped')} | | | |")
            continue
        out.append(f"| {r['model']} | {ct.get('compliance','—')}"
                   f" ({ct.get('verdict','—')}) | {th.get('themes','—')} | "
                   f"{th.get('unassigned','—')} | "
                   f"{th.get('dimensions_covered','—')}/{th.get('dimensions_total','—')}"
                   f" | {th.get('minutes','—')} |")
    return "\n".join(out)


def main(argv):
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--models", required=True,
                    help="逗號分隔的模型名稱")
    ap.add_argument("--provider", default=LLM.OLLAMA, choices=LLM.PROVIDERS)
    ap.add_argument("--base-url", default="")
    ap.add_argument("--api-key", default=os.environ.get("TACIT_API_KEY", ""))
    ap.add_argument("--num-ctx", type=int, default=LLM.DEFAULT_NUM_CTX)
    ap.add_argument("--framework", default="ri_stilgoe_2013")
    ap.add_argument("--corpus", default="en,zh",
                    help="要跑哪幾份語料，逗號分隔（en / zh）")
    ap.add_argument("--limit", type=int, default=None,
                    help="每份語料只跑前 N 篇（試跑用）")
    # 分窗參數要跟 app.py 的預設一致，否則量到的不是使用者會遇到的行為。
    # 開放調整是為了「窗口大小對召回的影響」這種比較，不是為了調校。
    ap.add_argument("--window", type=int, default=CH.DEFAULT_WINDOW_CHARS,
                    help=f"分窗長度（字元），預設 {CH.DEFAULT_WINDOW_CHARS}")
    ap.add_argument("--overlap", type=int, default=CH.DEFAULT_OVERLAP_CHARS,
                    help=f"重疊長度（字元），預設 {CH.DEFAULT_OVERLAP_CHARS}")
    # 雲端與本地可以同時跑，但必須寫到不同目錄，否則後結束的那一個會把
    # 先結束的結果檔整個覆蓋掉。
    ap.add_argument("--out-dir", default=OUT_DIR,
                    help=f"輸出目錄，預設 {OUT_DIR}")
    ap.add_argument("--probe-only", action="store_true",
                    help="只跑輸出約定探測，幾秒鐘就有結果")
    a = ap.parse_args(argv)

    out_dir = a.out_dir
    results_path = os.path.join(out_dir, "bench_results.json")
    os.makedirs(out_dir, exist_ok=True)
    # 批次腳本要比介面更有耐心（見 tacit_llm.set_retry_delays）
    LLM.set_retry_delays(5, 15, 40, 90)
    F.activate_by_id(a.framework)
    print(f"框架：{F.active().name()}　維度：{', '.join(F.active().dimensions)}")

    done = {}
    if os.path.isfile(results_path):
        with open(results_path, encoding="utf-8") as f:
            done = {r["model"]: r for r in json.load(f).get("results", [])}
        if done:
            print(f"（已有 {len(done)} 個模型的結果，會沿用；要重跑請刪掉 {results_path}）")

    results = []
    for model in [m.strip() for m in a.models.split(",") if m.strip()]:
        prev = done.get(model)
        prev_coding = (prev or {}).get("coding") or {}
        prev_themes = (prev or {}).get("themes") or {}
        prev_incomplete = prev and (
            any(c.get("failed") or c.get("stopped_on_quota") for c in prev_coding.values())
            or not prev_coding
            or prev_themes.get("error") or prev_themes.get("skipped"))
        if prev_incomplete:
            # 沒跑完的結果不能沿用：缺套件、額度用完、網路斷掉那一次會變成
            # 永久的表格內容。已完整編完的逐字稿會從存檔接續，不會重送。
            print(f"\n=== {model} — 上次沒有跑完，接著跑 ===")
        elif prev and not a.probe_only:
            print(f"\n=== {model} — 沿用既有結果 ===")
            results.append(done[model])
            continue
        print(f"\n=== {model} ===")
        ep = LLM.Endpoint(provider=a.provider, model=model, api_key=a.api_key,
                          base_url=a.base_url, num_ctx=a.num_ctx)
        row = {"model": model, "endpoint": ep.describe(),
               "provenance": ep.provenance(temperature=0.2, window_chars=a.window,
                                           overlap_chars=a.overlap),
               "framework": a.framework, "at": datetime.now().isoformat(timespec="seconds")}

        prev_contract = (prev or {}).get("contract") or {}
        if prev_contract.get("answered") and not prev_contract.get("error"):
            # 上次已經探測成功就不再花額度重測
            row["contract"] = prev_contract
            print("  輸出約定探測…沿用上次結果", end="", flush=True)
        else:
            print("  輸出約定探測…", end="", flush=True)
            row["contract"] = bench_contract(ep)
        c = row["contract"]
        print(f" {c['compliance']:.2f}（{c['verdict']}）{c['seconds']}s")
        if c.get("error"):
            print(f"      {str(c['error'])[:160]}")
        for nc in c["non_conforming"][:3]:
            print(f"      模型寫了：{nc['wrote']!r}")

        if not a.probe_only:
            row["coding"] = {}
            all_records = []
            for corpus in [c.strip() for c in a.corpus.split(",") if c.strip()]:
                lang = "zh-Hant" if corpus == "zh" else "en"
                print(f"  編碼 {corpus}…", flush=True)
                safe = re.sub(r"[^A-Za-z0-9._-]+", "_", model)
                stats, recs = bench_coding(
                    ep, corpus, lang, a.limit, a.window, a.overlap,
                    save_dir=os.path.join(out_dir, "records", safe, corpus))
                row["coding"][corpus] = stats
                all_records += recs
                print(f"  {corpus}：完成 {stats['completed']}/{stats['transcripts']} 份"
                      f"（其中 {stats['resumed']} 份沿用存檔）　"
                      f"{stats['segments']} 段（每萬字 {stats['per_10k']}）"
                      f" / {stats['codes']} 碼 / {stats['minutes']} 分")
                if stats["absence_codes"]:
                    print(f"      以缺席為證據而被擋下：{stats['absence_codes']} 碼")
                if stats["degenerate_titles"]:
                    print(f"      標題填成維度名：{stats['degenerate_titles']} 段")
            print("  主題歸納…", end="", flush=True)
            if any(s.get("stopped_on_quota") for s in row["coding"].values()):
                # 編碼沒跑完就做主題歸納，量到的是半批語料的主題，而且白花額度
                row["themes"] = {"skipped": "coding incomplete; rerun with another key"}
            else:
                row["themes"] = bench_themes(ep, all_records)
            th = row["themes"]
            if th.get("error") or th.get("skipped"):
                print(f" {th.get('error') or th.get('skipped')}")
            else:
                print(f" {th['themes']} 主題，未歸屬 {th['unassigned']}，"
                      f"涵蓋 {th['dimensions_covered']}/{th['dimensions_total']} 維度")
                if th["empty_dimensions"]:
                    print(f"      掛零的維度：{', '.join(th['empty_dimensions'])}")
                for u in th["unrecognised_dimensions"][:3]:
                    print(f"      不合格的維度值：{u}")
        results.append(row)

    with open(results_path, "w", encoding="utf-8") as f:
        json.dump({"generated": datetime.now().isoformat(timespec="seconds"),
                   "framework": a.framework, "results": results},
                  f, ensure_ascii=False, indent=2)
    md = markdown_table(results)
    with open(os.path.join(out_dir, "bench_table.md"), "w", encoding="utf-8") as f:
        f.write(md + "\n")
    print("\n" + "=" * 66)
    print(md)
    print("=" * 66)
    print(f"\n已寫入 {results_path} 與 {out_dir}/bench_table.md")
    print("參考編碼是固定比較基準，不是 ground truth——不要把一致性當成正確率。")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
